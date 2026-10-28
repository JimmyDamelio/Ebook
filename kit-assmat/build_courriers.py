#!/usr/bin/env python3
"""Builds dist/Modèles de courriers.docx (5 editable letters for the parent employer)."""
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Pt, RGBColor, Mm

TEAL = RGBColor(0x1F, 0x4D, 0x4A)
TERRA = RGBColor(0xC8, 0x64, 0x3F)
GREY = RGBColor(0x6E, 0x7C, 0x7B)

LETTERS = [
    ("Lettre de retrait de l'enfant",
     "À envoyer en recommandé avec accusé de réception, ou à remettre en main propre contre décharge "
     "(faites signer et dater un double). Le préavis démarre à la première présentation de la lettre.",
     [("left", "[Prénom Nom des parents]\n[Adresse]\n[Code postal, ville]"),
      ("right", "[Prénom Nom de l'assistante maternelle]\n[Adresse]\n[Code postal, ville]"),
      ("right", "À [ville], le [date]"),
      ("bold", "Objet : retrait de l'enfant [prénom] — fin du contrat de travail"),
      ("p", "Madame,"),
      ("p", "Par la présente, nous vous informons de notre décision de mettre fin au contrat de travail qui nous lie "
            "depuis le [date d'embauche] pour l'accueil de notre enfant [prénom], en exerçant notre droit de retrait "
            "conformément à l'article 119 de la convention collective nationale des particuliers employeurs et de "
            "l'emploi à domicile."),
      ("p", "Compte tenu de votre ancienneté de [nombre] mois, votre préavis est de [8 jours calendaires / 15 jours "
            "calendaires / 1 mois]. Il débute à la date de première présentation de ce courrier. Votre contrat "
            "prendra donc fin le [date]."),
      ("p", "À cette date, nous vous remettrons votre dernier bulletin de salaire, votre certificat de travail, le "
            "reçu pour solde de tout compte et l'attestation destinée à France Travail. Les sommes dues "
            "(congés payés, [régularisation de salaire,] [indemnité de rupture]) vous seront versées avec votre "
            "dernier salaire."),
      ("p", "Nous vous remercions pour [la qualité de l'accueil réservé à notre enfant] et vous prions d'agréer, "
            "Madame, l'expression de nos salutations distinguées."),
      ("right", "[Signatures des parents]")]),
    ("Certificat de travail",
     "À remettre le dernier jour du contrat (article 69 de la convention collective).",
     [("bold", "CERTIFICAT DE TRAVAIL"),
      ("p", "Nous soussignés, [Prénom Nom des parents], demeurant [adresse complète], particuliers employeurs, "
            "certifions avoir employé :"),
      ("p", "Madame [Prénom Nom], demeurant [adresse complète],"),
      ("p", "en qualité d'assistante maternelle agréée, du [date d'embauche] au [date de fin du contrat], pour "
            "l'accueil de notre enfant [prénom]."),
      ("p", "Madame [Nom] bénéficie du maintien à titre gratuit des garanties de prévoyance (portabilité) auprès de "
            "l'IRCEM Prévoyance, dans les conditions prévues par la convention collective."),
      ("p", "Madame [Nom] nous quitte libre de tout engagement."),
      ("right", "Fait à [ville], le [date]\n[Signatures des parents]")]),
    ("Reçu pour solde de tout compte",
     "Établi en deux exemplaires. Il détaille les sommes versées à la fin du contrat ; les montants sont dans "
     "l'onglet « Fin de contrat » du classeur.",
     [("bold", "REÇU POUR SOLDE DE TOUT COMPTE"),
      ("p", "Je soussignée, Madame [Prénom Nom], assistante maternelle agréée, demeurant [adresse], reconnais avoir "
            "reçu de [Prénom Nom des parents], demeurant [adresse], pour solde de tout compte, la somme de "
            "[montant] € net, versée le [date] par [virement / chèque / Pajemploi+], en paiement des salaires, "
            "accessoires du salaire et indemnités de toute nature dus au titre de l'exécution et de la cessation "
            "de mon contrat de travail, qui a pris fin le [date]."),
      ("p", "Détail des sommes :"),
      ("li", "Salaire du mois de [mois] : [montant] € brut"),
      ("li", "Régularisation de salaire (accueil sur 46 semaines ou moins) : [montant] € brut"),
      ("li", "Indemnité compensatrice de congés payés : [montant] € brut"),
      ("li", "Indemnité de rupture (non soumise à cotisations) : [montant] €"),
      ("li", "Indemnités d'entretien et de repas du mois : [montant] €"),
      ("p", "Ce reçu est établi en deux exemplaires, dont un m'est remis. Il peut être dénoncé dans les six mois qui "
            "suivent sa signature, délai au-delà duquel il devient libératoire pour l'employeur pour les sommes qui "
            "y sont mentionnées."),
      ("right", "Fait à [ville], le [date]\nSignature de l'assistante maternelle\n(précédée de la mention "
                "« Pour solde de tout compte »)")]),
    ("Régularisation prévisionnelle annuelle",
     "Accueil sur 46 semaines ou moins : à faire à chaque date anniversaire du contrat. Aucun paiement à cette "
     "date : les régularisations annuelles se compensent et sont soldées à la fin du contrat.",
     [("bold", "RÉGULARISATION PRÉVISIONNELLE — ANNÉE DU [date] AU [date]"),
      ("p", "Entre [Prénom Nom des parents], particuliers employeurs, et Madame [Prénom Nom], assistante "
            "maternelle agréée, pour l'accueil de [prénom de l'enfant], contrat du [date d'embauche] — accueil sur "
            "[nombre] semaines par an."),
      ("li", "Heures de travail réellement effectuées sur la période (hors heures complémentaires et majorées, "
             "jours fériés chômés et absences payées compris) : [nombre] h"),
      ("li", "Heures payées au titre du salaire mensualisé sur la période : [nombre] h"),
      ("li", "Écart : [nombre] h, soit [montant] € brut au taux horaire de [taux] €, "
             "[en faveur de l'assistante maternelle / en faveur des parents employeurs]."),
      ("p", "Les parties prennent acte de ce calcul. Il ne donne lieu à aucun versement à ce jour. Il sera repris, "
            "avec les régularisations des autres années, dans la régularisation définitive effectuée à la fin du "
            "contrat (article 124 de la convention collective) ; seul un solde en faveur de l'assistante "
            "maternelle donne alors lieu à paiement."),
      ("p", "Planning des semaines d'accueil programmées pour les 12 prochains mois : [semaines ou périodes]."),
      ("right", "Fait en deux exemplaires à [ville], le [date]\nSignature des parents        "
                "Signature de l'assistante maternelle")]),
    ("Dates de congés payés",
     "Avec un seul employeur, c'est vous qui fixez les dates et les communiquez au plus tard le 1er mars "
     "(article 102-1-1 de la convention collective).",
     [("right", "À [ville], le [date — au plus tard le 1er mars]"),
      ("bold", "Objet : dates des congés payés [année]"),
      ("p", "Madame,"),
      ("p", "Conformément à la convention collective, nous vous communiquons les dates de vos congés payés pour "
            "l'année [année] :"),
      ("li", "Congé principal : du [date] au [date] inclus ([nombre] semaines, entre le 1er mai et le 31 octobre) ;"),
      ("li", "[Autre période] : du [date] au [date] inclus."),
      ("p", "Vous avez acquis [nombre] jours ouvrables de congés payés au 31 mai [année]. [Les jours de repos "
            "au-delà de vos droits acquis seront pris en congés sans solde.]"),
      ("p", "Nous restons disponibles pour en parler et vous prions d'agréer, Madame, l'expression de nos "
            "salutations distinguées."),
      ("right", "[Signatures des parents]\nReçu le [date] — signature de l'assistante maternelle")]),
]


def build(path):
    doc = Document()
    for s in doc.sections:
        s.page_height, s.page_width = Mm(297), Mm(210)
        s.left_margin = s.right_margin = Mm(22)
        s.top_margin, s.bottom_margin = Mm(20), Mm(20)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(11)
    st.paragraph_format.space_after = Pt(8)
    st.paragraph_format.line_spacing = 1.15

    t = doc.add_paragraph()
    r = t.add_run("NID SEREIN · OUTILS POUR PARENTS EMPLOYEURS")
    r.font.size, r.font.bold, r.font.color.rgb = Pt(9), True, TERRA
    h = doc.add_paragraph()
    r = h.add_run("Modèles de courriers — Assistante maternelle")
    r.font.size, r.font.bold, r.font.color.rgb = Pt(22), True, TEAL
    intro = doc.add_paragraph("Cinq modèles à compléter : remplacez les passages entre [crochets] et supprimez les "
                              "mentions qui ne vous concernent pas. Ces modèles sont fournis à titre indicatif et ne "
                              "remplacent pas un conseil juridique.")
    intro.runs[0].font.color.rgb = GREY
    for i, (title, note, _) in enumerate(LETTERS, 1):
        p = doc.add_paragraph(f"{i}.  {title}")
        p.runs[0].font.bold, p.runs[0].font.color.rgb = True, TEAL

    for title, note, body in LETTERS:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        p = doc.add_paragraph()
        r = p.add_run(title.upper())
        r.font.size, r.font.bold, r.font.color.rgb = Pt(9), True, TERRA
        n = doc.add_paragraph(note)
        n.runs[0].font.size, n.runs[0].font.italic, n.runs[0].font.color.rgb = Pt(9), True, GREY
        doc.add_paragraph()
        for kind, text in body:
            if kind == "li":
                p = doc.add_paragraph(text, style="List Bullet")
            else:
                p = doc.add_paragraph()
                for j, line in enumerate(text.split("\n")):
                    run = p.add_run(line)
                    if j < len(text.split("\n")) - 1:
                        run.add_break()
                    if kind == "bold":
                        run.font.bold = True
                        run.font.color.rgb = TEAL
                if kind == "right":
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                if kind == "p":
                    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    doc.save(path)
    return path


if __name__ == "__main__":
    here = os.path.dirname(os.path.abspath(__file__))
    print(build(os.path.join(here, "dist", "Modèles de courriers.docx")))
