"""Recalculate an xlsx with headless LibreOffice and return an openpyxl workbook of cached values."""
import os, subprocess, tempfile, shutil
from openpyxl import load_workbook


def recalc(path):
    out = tempfile.mkdtemp()
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML",
                    "--outdir", out, path], check=True, capture_output=True, timeout=300)
    res = os.path.join(out, os.path.basename(path))
    wb = load_workbook(res, data_only=True)
    shutil.rmtree(out)
    return wb


def errors(wb):
    bad = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and (c.value.startswith("#") or c.value.startswith("Err:")):
                    bad.append((ws.title, c.coordinate, c.value))
    return bad
